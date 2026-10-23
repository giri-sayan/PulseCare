"""
Master Report Generator for BCSE302P Database Systems Laboratory
Generates the complete project document based on Dr. G. M. Karthik's template.
Student: Sayan Giri (24BCT0281)
"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

TEMPLATE_PATH = r"C:\Users\USER\Downloads\BCSE302P_Mini_Project_Report_Template.docx"
OUTPUT_DOCX_PATH = r"C:\Users\USER\Desktop\DBMS\report\BCSE302P_Project_24BCT0281.docx"

doc = docx.Document(TEMPLATE_PATH)

# Helper functions
def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def style_run(r, font_name="Times New Roman", font_size=12, bold=False, italic=False, color=(0,0,0)):
    r.font.name = font_name
    r.font.size = Pt(font_size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = RGBColor(*color)

def set_p_formatting(p, line_spacing=1.5, space_after=6, align=WD_ALIGN_PARAGRAPH.LEFT):
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_after = Pt(space_after)
    p.alignment = align

print(f"Loaded template with {len(doc.paragraphs)} paragraphs and {len(doc.tables)} tables.")

# =============================================================================
# 1. POPULATE COVER PAGE & CERTIFICATE
# =============================================================================
for p in doc.paragraphs[:35]:
    if "<Project Title>" in p.text:
        p.text = p.text.replace("<Project Title>", "PulseCare: Smart Hospital Resource, Inpatient Admission & Clinical Billing Management System")
        set_p_formatting(p, align=WD_ALIGN_PARAGRAPH.CENTER)
        for r in p.runs:
            style_run(r, font_size=15, bold=True, color=(15, 23, 42))
    elif "Date of Submission:" in p.text:
        p.text = "Date of Submission: 23 / 10 / 2026"
        set_p_formatting(p, align=WD_ALIGN_PARAGRAPH.CENTER)
        for r in p.runs:
            style_run(r, font_size=12, bold=True)
    elif "<Name 1" in p.text:
        p.text = "This is to certify that the Mini Project Report titled \"PulseCare: Smart Hospital Resource, Inpatient Admission & Clinical Billing Management System\" submitted by Sayan Giri (Reg. No. 24BCT0281) in partial fulfilment of the requirements of the course BCSE302P – Database Systems Laboratory, is a record of original work carried out under my supervision during the Fall Semester 2026–27."
        set_p_formatting(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        for r in p.runs:
            style_run(r, font_size=12)
    elif "Date:" in p.text and "Place:" not in p.text:
        p.text = "Date: 23 / 10 / 2026"
        for r in p.runs:
            style_run(r, font_size=12)

# Update Table 0 (Team details) - Keep row 0 (header) and row 1 (Sayan Giri), remove rows 2, 3, 4
t0 = doc.tables[0]
for r in list(t0.rows)[2:]:
    t0._element.remove(r._element)

t0.rows[1].cells[0].text = "1"
t0.rows[1].cells[1].text = "Sayan Giri"
t0.rows[1].cells[2].text = "24BCT0281"

for row in t0.rows:
    for c in row.cells:
        set_cell_margins(c, 80, 80, 100, 100)
        p = c.paragraphs[0]
        set_p_formatting(p, line_spacing=1.15, space_after=2)
        for r in p.runs:
            style_run(r, font_size=11, bold=(row == t0.rows[0]))

print("Cover page & Certificate updated.")

# =============================================================================
# 2. ABSTRACT
# =============================================================================
abstract_text = (
    "Modern healthcare institutions demand high-integrity, concurrent, and automated data management systems "
    "to eliminate clinical errors and administrative bottlenecks. This project presents PulseCare, an enterprise-grade "
    "hospital resource, inpatient admission, and clinical billing management system designed and implemented on "
    "Oracle Database 21c Express Edition. The database architecture comprises 12 normalized relational entities structured "
    "in Boyce-Codd Normal Form (BCNF) to guarantee data integrity, eliminate redundancy, and maintain referential consistency. "
    "Advanced Oracle PL/SQL programming constructs—including stored functions, stored procedures with transactional controls, "
    "an encapsulated hospital billing package, database triggers (incorporating a compound trigger to eliminate mutating-table "
    "anomalies), and parameterized explicit cursors—are developed to enforce complex hospital policies such as bed state "
    "synchronization, appointment collision prevention, and tiered insurance billing calculations. A responsive web-based "
    "front-end built with Python and Streamlit integrates directly with the Oracle pluggable database via the python-oracledb "
    "connector. Rigorous verification across 11 comprehensive test cases demonstrates complete adherence to clinical operational "
    "constraints, automated financial auditing, and real-time bed utilization tracking, establishing a dependable software "
    "foundation for contemporary hospital administration."
)

for p in doc.paragraphs[35:40]:
    if "Write 150" in p.text:
        p.text = abstract_text
        set_p_formatting(p, line_spacing=1.5, space_after=8, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        for r in p.runs:
            style_run(r, font_size=12)
    elif "Keywords:" in p.text:
        p.text = "Keywords: DBMS, Oracle Database 21c XE, PL/SQL, Hospital Management System, BCNF Normalization, Compound Triggers, Streamlit"
        set_p_formatting(p, line_spacing=1.5, space_after=12)
        for r in p.runs:
            style_run(r, font_size=12, italic=True)

print("Abstract updated.")

# Save intermediate document to check progress
doc.save(OUTPUT_DOCX_PATH)
print(f"Saved initial draft to {OUTPUT_DOCX_PATH}")
